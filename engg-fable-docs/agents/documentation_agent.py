"""agents/documentation_agent.py — Professional manual assembly (DOCX + PDF).

Document structure per master.md §4:
  Cover page → Revision history + TOC → BOM → System overview (block diagram +
  system description) → per-subsystem sections (diagram, multi-level description,
  interconnection table, signal list) → connectivity appendix.
"""
import os
from datetime import date
from typing import Dict

import pandas as pd
from PIL import Image as PILImage

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (BaseDocTemplate, Frame, Image as RLImage,
                                PageBreak, PageTemplate, Paragraph, Spacer,
                                Table, TableStyle)
from reportlab.platypus.tableofcontents import TableOfContents

from src.config import DOC_AUTHOR, DOC_NUMBER, DOC_ORGANIZATION, DOC_VERSION, OUTPUT_DIR

ACCENT = "#1F4E79"


# ══════════════════════════ DOCX helpers ══════════════════════════════════
def _docx_toc_field(doc):
    """Insert a Word TOC field (populated when the user updates fields in Word)."""
    para = doc.add_paragraph()
    run = para.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    hint = OxmlElement("w:t")
    hint.text = "Right-click here and choose 'Update Field' to generate the Table of Contents."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, sep, hint, end):
        run._r.append(el)


def _docx_table(doc, header, rows, style="Light Grid Accent 1"):
    t = doc.add_table(rows=1, cols=len(header))
    try:
        t.style = style
    except KeyError:
        t.style = "Table Grid"
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = str(h)
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    return t


# ══════════════════════════ PDF helpers ═══════════════════════════════════
class _ManualDoc(BaseDocTemplate):
    """BaseDocTemplate that feeds headings into the TableOfContents flowable."""

    def __init__(self, filename, sys_title="", **kw):
        super().__init__(filename, pagesize=letter, **kw)
        self._sys_title = sys_title
        frame = Frame(0.9 * inch, 0.9 * inch, letter[0] - 1.8 * inch,
                      letter[1] - 1.8 * inch, id="body")
        self.addPageTemplates([PageTemplate(id="main", frames=[frame],
                                            onPage=self._footer)])

    def _footer(self, canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.grey)
        canvas.drawString(0.9 * inch, 0.55 * inch,
                          f"{self._sys_title} — Technical User Manual  |  {DOC_NUMBER} Rev {DOC_VERSION}")
        canvas.drawRightString(letter[0] - 0.9 * inch, 0.55 * inch, f"Page {doc.page}")
        canvas.restoreState()

    def afterFlowable(self, flowable):
        if isinstance(flowable, Paragraph):
            style = flowable.style.name
            if style == "TOCHeading1":
                self.notify("TOCEntry", (0, flowable.getPlainText(), self.page))
            elif style == "TOCHeading2":
                self.notify("TOCEntry", (1, flowable.getPlainText(), self.page))


def _pdf_styles():
    ss = getSampleStyleSheet()
    ss.add(ParagraphStyle("CoverTitle", parent=ss["Title"], fontSize=30, leading=36,
                          textColor=colors.HexColor(ACCENT), spaceAfter=18))
    ss.add(ParagraphStyle("CoverSub", parent=ss["Normal"], fontSize=16, leading=20,
                          alignment=1, textColor=colors.grey))
    ss.add(ParagraphStyle("TOCHeading1", parent=ss["Heading1"],
                          textColor=colors.HexColor(ACCENT), spaceBefore=14))
    ss.add(ParagraphStyle("TOCHeading2", parent=ss["Heading2"],
                          textColor=colors.HexColor(ACCENT), spaceBefore=10))
    ss.add(ParagraphStyle("Cell", parent=ss["Normal"], fontSize=7.5, leading=9))
    ss.add(ParagraphStyle("CellHead", parent=ss["Normal"], fontSize=8, leading=10,
                          textColor=colors.white, fontName="Helvetica-Bold"))
    ss.add(ParagraphStyle("Body", parent=ss["Normal"], fontSize=9.5, leading=13,
                          spaceAfter=6))
    return ss


def _pdf_table(header, rows, ss, col_widths=None):
    data = [[Paragraph(str(h), ss["CellHead"]) for h in header]]
    for row in rows:
        data.append([Paragraph(str(v).replace("→", "&rarr;"), ss["Cell"]) for v in row])
    t = Table(data, repeatRows=1, colWidths=col_widths, hAlign="LEFT")
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(ACCENT)),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F6FA")]),
    ]))
    return t


def _pdf_image(path, max_w=6.4 * inch, max_h=7.5 * inch):
    try:
        with PILImage.open(path) as img:
            w, h = img.size
        scale = min(max_w / w, max_h / h, 1.0)
        return RLImage(path, width=w * scale, height=h * scale)
    except Exception:
        return None


# ══════════════════════════ Agent ═════════════════════════════════════════
class DocumentationAgent:

    @staticmethod
    def run(state: dict, output_dir: str = OUTPUT_DIR) -> Dict[str, str]:
        os.makedirs(output_dir, exist_ok=True)
        docx_path = os.path.join(output_dir, "Technical_User_Manual.docx")
        pdf_path = os.path.join(output_dir, "Technical_User_Manual.pdf")

        df_conn = state.get("df_connectivity", pd.DataFrame())
        sys_name = (str(df_conn["System_Name"].iloc[0])
                    if "System_Name" in df_conn.columns and not df_conn.empty else "System")
        sys_title = sys_name.replace("_", " ")

        DocumentationAgent._build_docx(state, sys_title, docx_path)
        DocumentationAgent._build_pdf(state, sys_title, pdf_path)
        print(f"  ✓ Documents saved: {docx_path}, {pdf_path}")
        return {"docx": docx_path, "pdf": pdf_path}

    # ── DOCX ──────────────────────────────────────────────────────────────
    @staticmethod
    def _build_docx(state, sys_title, path):
        df_conn = state.get("df_connectivity", pd.DataFrame())
        df_bom = state.get("df_bom", pd.DataFrame())
        descs = state.get("descriptions", {})
        diags = state.get("diagrams", {})
        interconnects = state.get("interconnections", {})
        system_desc = state.get("system_description", "")
        today = date.today().isoformat()

        doc = Document()
        doc.core_properties.title = f"{sys_title} — Technical User Manual"
        doc.core_properties.author = DOC_AUTHOR

        # ── Cover page ──
        for _ in range(6):
            doc.add_paragraph()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(sys_title)
        r.font.size = Pt(36)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("Technical User Manual")
        r.font.size = Pt(20)
        r.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
        for _ in range(8):
            doc.add_paragraph()
        meta = doc.add_paragraph()
        meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        meta.add_run(
            f"Document No: {DOC_NUMBER}    Version: {DOC_VERSION}    Date: {today}\n"
            f"Prepared by: {DOC_AUTHOR}"
            + (f"\n{DOC_ORGANIZATION}" if DOC_ORGANIZATION else "")
        ).font.size = Pt(11)
        doc.add_page_break()

        # ── Revision history + TOC ──
        doc.add_heading("Revision History", 1)
        _docx_table(doc, ["Version", "Date", "Author", "Description"],
                    [[DOC_VERSION, today, DOC_AUTHOR, "Initial release"]])
        doc.add_paragraph()
        doc.add_heading("Table of Contents", 1)
        _docx_toc_field(doc)
        doc.add_page_break()

        # ── 1. BOM ──
        doc.add_heading("1. Bill of Materials", 1)
        if not df_bom.empty:
            _docx_table(doc, list(df_bom.columns),
                        [[r[c] for c in df_bom.columns] for _, r in df_bom.iterrows()])
        doc.add_page_break()

        # ── 2. System overview ──
        doc.add_heading("2. System Overview", 1)
        sys_diag = diags.get("System_Overview")
        if sys_diag and os.path.exists(sys_diag):
            doc.add_picture(sys_diag, width=Inches(6.3))
        doc.add_heading("2.1 Functional Description", 2)
        for para in (system_desc or "System description not available.").split("\n\n"):
            doc.add_paragraph(para)

        bridges = state.get("bridges", [])
        if bridges:
            doc.add_heading("2.2 Subsystem Interconnections", 2)
            _docx_table(doc, ["Bridging Component", "Subsystems", "Signal Classes"],
                        [[b["component"], " ↔ ".join(b["subsystems"]),
                          ", ".join(b["signal_classes"])] for b in bridges])
        doc.add_page_break()

        # ── 3. Subsystems ──
        doc.add_heading("3. Subsystem Descriptions", 1)
        for idx, (sn, d) in enumerate(descs.items(), start=1):
            if not isinstance(d, dict):
                continue
            title = sn.replace("_", " ")
            doc.add_heading(f"3.{idx} {title}", 2)

            ip = diags.get(sn)
            if ip and os.path.exists(ip):
                doc.add_picture(ip, width=Inches(6.3))

            doc.add_heading(f"3.{idx}.1 Functional Description", 3)
            for para in d.get("full_description", "").split("\n\n"):
                doc.add_paragraph(para)

            comp_descs = d.get("component_descriptions", {})
            if comp_descs:
                doc.add_heading(f"3.{idx}.2 Component Roles", 3)
                for text in comp_descs.values():
                    doc.add_paragraph(text, style="List Bullet")

            ic = interconnects.get(sn, [])
            if ic:
                doc.add_heading(f"3.{idx}.3 Interconnections with Other Subsystems", 3)
                _docx_table(doc, ["Connected Subsystem", "Via Components", "Signals"],
                            [[r["other_subsystem"].replace("_", " "),
                              r["shared_components"], r["signals"]] for r in ic])

            sigs = d.get("signals", [])
            if sigs:
                doc.add_heading(f"3.{idx}.4 Signal List", 3)
                _docx_table(doc, ["Signal", "Class", "Connections"],
                            [[s["signal"], s["class"], s["connections"]] for s in sigs])
            doc.add_page_break()

        # ── 4. Connectivity appendix ──
        doc.add_heading("4. Appendix — Connectivity Data", 1)
        for sn, g in df_conn.groupby("Subsystem_Name"):
            doc.add_heading(str(sn).replace("_", " "), 2)
            cols = ["Component_ID", "Source_Pin", "Target_ID", "Target_Pin", "Signal_Name"]
            cols = [c for c in cols if c in g.columns]
            _docx_table(doc, cols, [[r[c] for c in cols] for _, r in g.iterrows()])

        doc.save(path)

    # ── PDF ───────────────────────────────────────────────────────────────
    @staticmethod
    def _build_pdf(state, sys_title, path):
        df_conn = state.get("df_connectivity", pd.DataFrame())
        df_bom = state.get("df_bom", pd.DataFrame())
        descs = state.get("descriptions", {})
        diags = state.get("diagrams", {})
        interconnects = state.get("interconnections", {})
        system_desc = state.get("system_description", "")
        bridges = state.get("bridges", [])
        today = date.today().isoformat()

        ss = _pdf_styles()
        doc = _ManualDoc(path, sys_title=sys_title)
        story = []

        # ── Cover ──
        story.append(Spacer(1, 2.2 * inch))
        story.append(Paragraph(sys_title, ss["CoverTitle"]))
        story.append(Paragraph("Technical User Manual", ss["CoverSub"]))
        story.append(Spacer(1, 2.4 * inch))
        story.append(_pdf_table(
            ["Document No", "Version", "Date", "Prepared By"],
            [[DOC_NUMBER, DOC_VERSION, today, DOC_AUTHOR]], ss))
        story.append(PageBreak())

        # ── Revision history + TOC ──
        story.append(Paragraph("Revision History", ss["TOCHeading1"]))
        story.append(_pdf_table(["Version", "Date", "Author", "Description"],
                                [[DOC_VERSION, today, DOC_AUTHOR, "Initial release"]], ss))
        story.append(Spacer(1, 0.35 * inch))
        story.append(Paragraph("Table of Contents", ss["TOCHeading1"]))
        toc = TableOfContents()
        toc.levelStyles = [
            ParagraphStyle("TOC1", fontSize=10, leading=14, leftIndent=6),
            ParagraphStyle("TOC2", fontSize=9, leading=12, leftIndent=20),
        ]
        story.append(toc)
        story.append(PageBreak())

        # ── 1. BOM ──
        story.append(Paragraph("1. Bill of Materials", ss["TOCHeading1"]))
        if not df_bom.empty:
            story.append(_pdf_table(list(df_bom.columns),
                                    [[r[c] for c in df_bom.columns]
                                     for _, r in df_bom.iterrows()], ss))
        story.append(PageBreak())

        # ── 2. System overview ──
        story.append(Paragraph("2. System Overview", ss["TOCHeading1"]))
        sys_diag = diags.get("System_Overview")
        if sys_diag and os.path.exists(sys_diag):
            img = _pdf_image(sys_diag)
            if img:
                story.append(img)
                story.append(Spacer(1, 0.15 * inch))
        for para in (system_desc or "").split("\n\n"):
            story.append(Paragraph(para, ss["Body"]))
        if bridges:
            story.append(Spacer(1, 0.15 * inch))
            story.append(Paragraph("Subsystem Interconnections", ss["TOCHeading2"]))
            story.append(_pdf_table(
                ["Bridging Component", "Subsystems", "Signal Classes"],
                [[b["component"], " / ".join(b["subsystems"]),
                  ", ".join(b["signal_classes"])] for b in bridges], ss))
        story.append(PageBreak())

        # ── 3. Subsystems ──
        story.append(Paragraph("3. Subsystem Descriptions", ss["TOCHeading1"]))
        for idx, (sn, d) in enumerate(descs.items(), start=1):
            if not isinstance(d, dict):
                continue
            title = sn.replace("_", " ")
            story.append(Paragraph(f"3.{idx} {title}", ss["TOCHeading2"]))
            ip = diags.get(sn)
            if ip and os.path.exists(ip):
                img = _pdf_image(ip)
                if img:
                    story.append(img)
                    story.append(Spacer(1, 0.12 * inch))
            for para in d.get("full_description", "").split("\n\n"):
                story.append(Paragraph(para, ss["Body"]))

            comp_descs = d.get("component_descriptions", {})
            if comp_descs:
                story.append(Paragraph("<b>Component Roles</b>", ss["Body"]))
                for text in comp_descs.values():
                    story.append(Paragraph(f"• {text}", ss["Body"]))

            ic = interconnects.get(sn, [])
            if ic:
                story.append(Spacer(1, 0.1 * inch))
                story.append(Paragraph("<b>Interconnections with Other Subsystems</b>", ss["Body"]))
                story.append(_pdf_table(
                    ["Connected Subsystem", "Via Components", "Signals"],
                    [[r["other_subsystem"].replace("_", " "), r["shared_components"],
                      r["signals"]] for r in ic], ss))

            sigs = d.get("signals", [])
            if sigs:
                story.append(Spacer(1, 0.1 * inch))
                story.append(Paragraph("<b>Signal List</b>", ss["Body"]))
                story.append(_pdf_table(["Signal", "Class", "Connections"],
                                        [[s["signal"], s["class"], s["connections"]]
                                         for s in sigs], ss,
                                        col_widths=[1.1 * inch, 1.2 * inch, 4.1 * inch]))
            story.append(PageBreak())

        # ── 4. Appendix ──
        story.append(Paragraph("4. Appendix — Connectivity Data", ss["TOCHeading1"]))
        for sn, g in df_conn.groupby("Subsystem_Name"):
            story.append(Paragraph(str(sn).replace("_", " "), ss["TOCHeading2"]))
            cols = ["Component_ID", "Source_Pin", "Target_ID", "Target_Pin", "Signal_Name"]
            cols = [c for c in cols if c in g.columns]
            story.append(_pdf_table(cols, [[r[c] for c in cols]
                                           for _, r in g.iterrows()], ss))
            story.append(Spacer(1, 0.2 * inch))

        doc.multiBuild(story)
