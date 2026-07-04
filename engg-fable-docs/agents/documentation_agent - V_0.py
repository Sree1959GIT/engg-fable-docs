"""agents/documentation_agent.py — Fixed PDF image handling, no preserveAspectRatio."""
import os, pandas as pd
from typing import Dict
from docx import Document
from docx.shared import Inches
from PIL import Image as PILImage  # Get actual dimensions
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle
from reportlab.platypus import Image as RLImage



class DocumentationAgent:
    @staticmethod
    def run(state: dict, output_dir: str = "output") -> Dict[str, str]:
        os.makedirs(output_dir, exist_ok=True)
        docx_path = os.path.join(output_dir, "Technical_User_Manual.docx")
        pdf_path = os.path.join(output_dir, "Technical_User_Manual.pdf")

        df_conn = state.get("df_connectivity", pd.DataFrame())
        df_bom = state.get("df_bom", pd.DataFrame())
        descs = state.get("descriptions", {})
        diags = state.get("diagrams", {})
        sys_name = df_conn["System_Name"].iloc[0] if "System_Name" in df_conn.columns else "System"

        # ══════════ WORD (.docx) ══════════
        doc = Document()
        doc.add_heading(f"{sys_name} - Technical User Manual", 0)
        doc.add_heading("1. Bill of Materials", 1)
        if not df_bom.empty:
            t = doc.add_table(rows=1, cols=len(df_bom.columns))
            t.style = "Light Grid Accent 1"
            for i, c in enumerate(df_bom.columns):
                t.rows[0].cells[i].text = str(c)
            for _, r in df_bom.iterrows():
                cells = t.add_row().cells
                for i, c in enumerate(df_bom.columns):
                    cells[i].text = str(r[c])

        doc.add_heading("2. System Overview", 1)
        for _, d in descs.items():
            if isinstance(d, dict):
                doc.add_paragraph(d.get("full_description", ""))
                break

        doc.add_heading("3. Subsystems & Wiring Diagrams", 1)
        for sn, d in descs.items():
            if isinstance(d, dict):
                doc.add_heading(sn, 2)
                doc.add_paragraph(d.get("full_description", ""))
                ip = diags.get(sn)
                if ip and os.path.exists(ip):
                    doc.add_picture(ip, width=Inches(5.5))

        doc.add_heading("4. Connectivity Data", 1)
        for sn, g in df_conn.groupby("Subsystem_Name"):
            doc.add_heading(sn, 2)
            t = doc.add_table(rows=1, cols=len(g.columns))
            t.style = "Light Grid Accent 1"
            for i, c in enumerate(g.columns):
                t.rows[0].cells[i].text = str(c)
            for _, r in g.iterrows():
                cells = t.add_row().cells
                for i, c in enumerate(g.columns):
                    cells[i].text = str(r[c])
        doc.save(docx_path)

        # ══════════ PDF ══════════
        pdf = SimpleDocTemplate(pdf_path, pagesize=letter)
        styles = getSampleStyleSheet()
        story = []
        story.append(Paragraph(f"{sys_name} - Technical User Manual",
                               styles["Title"]))
        story.append(Spacer(1, 0.5 * inch))

        if not df_bom.empty:
            story.append(Paragraph("1. Bill of Materials", styles["Heading1"]))
            pdf_data = [list(df_bom.columns)]
            for _, r in df_bom.iterrows():
                pdf_data.append([str(r[c]) for c in df_bom.columns])
            t = Table(pdf_data, repeatRows=1)
            t.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1f4e79")),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
            ]))
            story.append(t)
            story.append(PageBreak())

        story.append(Paragraph("2. Subsystem Descriptions & Diagrams",
                               styles["Heading1"]))
        for sn, d in descs.items():
            if isinstance(d, dict):
                story.append(Paragraph(sn, styles["Heading2"]))
                story.append(Paragraph(d.get("full_description", ""),
                                       styles["Normal"]))
                ip = diags.get(sn)
                if ip and os.path.exists(ip):
                    # Open with PIL to get aspect ratio, then use width only
                    try:
                        with PILImage.open(ip) as img:
                            w, h = img.size
                            max_w = 450
                            ratio = max_w / w
                            display_h = h * ratio
                        story.append(RLImage(ip, width=max_w, height=display_h))
                    except Exception:
                        story.append(RLImage(ip, width=400))
                story.append(PageBreak())

        pdf.build(story)
        print(f"  ✓ Documents saved: {docx_path}, {pdf_path}")
        return {"docx": docx_path, "pdf": pdf_path}
